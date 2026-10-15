# build_full_document.py
# Generación del informe formal del Taller 10 para Gianfranco Arocutipa
# Basado en la plantilla SI886-PLANTILLA-TALLER (3).docx

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPO_URL = "https://github.com/GianfrancoArocutipa/Taller10_ArocutipaGianfranco/tree/taller-10"

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table_header(row, col_widths, bg_color='1F4E79'):
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        cell.width = col_widths[i]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = 'Calibri'

def add_styled_row(table, data, col_widths, bg_color='FFFFFF', bold_col=None, align_center_cols=[]):
    row = table.add_row()
    for i, (cell, val) in enumerate(zip(row.cells, data)):
        if bg_color != 'FFFFFF':
            set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        cell.width = col_widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in align_center_cols else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(val))
        run.font.name = 'Calibri'
        run.font.size = Pt(8.5)
        if bold_col is not None and i == bold_col:
            run.font.bold = True
    return row

print("Iniciando generación de documento...")
doc = Document('SI886-PLANTILLA-TALLER (3).docx')

# 1. Modificar Carátula
for p in doc.paragraphs:
    if "“…………………………………………….”" in p.text:
        p.text = "“TALLER 10: MODELADO DE LA ARQUITECTURA ACTUAL Y OBJETIVO EN ARCHI”\n(CASO: AGROEXPORTADORA VALLE VIEJO S.A.C.)"
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
    elif "“………………………….”" in p.text:
        p.text = "“PLANEAMIENTO ESTRATÉGICO DE TECNOLOGÍAS DE LA INFORMACIÓN”"
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
    elif "APELLIDOS Y NOMBRES · CÓDIGO" in p.text:
        p.text = "AROCUTIPA AROCUTIPA, GIANFRANCO · 2023076790"
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
    elif "DR. OSCAR JUAN JIMENEZ FLORES" in p.text or "DR. OSCAR JUAN JIMENEZ FLORES" in p.text:
        p.text = "DR. OSCAR JUAN JIMENEZ FLORES"
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
    elif "TACNA — PERÚ" in p.text:
        p.text = "TACNA — PERÚ\n2026"
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)

print("Carátula actualizada.")
doc.save("test_caratula.docx")
print("Prueba carátula guardada.")
